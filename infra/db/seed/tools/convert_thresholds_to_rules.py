#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convertit gabarit_seuils_cliniques.xlsx (une fois rempli par le concepteur
médical) en un script SQL de propositions de clinical_rules — TOUJOURS
`active = false`, jamais d'activation automatique (§57, §59, §78).

Usage :
    python3 convert_thresholds_to_rules.py gabarit_seuils_cliniques.xlsx > proposed_thresholds.sql

Le SQL généré doit être relu ligne par ligne par le concepteur médical avant
tout `active = true`.

Types de colonne « Type » reconnus :
  - boolean    : condition {"field": code, "operator": "equals", "value": true}
  - scale_0_10 : condition {"field": code, "operator": "gte", "value": <seuil>}
  - text       : aucune condition dérivable automatiquement (à écrire à la main)
  - select_3   : AJOUTÉ le 21/08/2026 (réponse Q1 au document de questions
                 ouvertes du Sprint 16) — item qualitatif à 3 niveaux nommés
                 (0 = vert, 1 = orange, 2 = rouge), même sémantique que le
                 type `select` de packages/domain/src/screening.ts. Génère
                 deux conditions `gte` (comme scale_0_10), avec un seuil par
                 défaut de 1 (orange) et 2 (rouge) si les colonnes « Seuil
                 orange »/« Seuil rouge » sont laissées vides — un item à 3
                 niveaux nommés n'a normalement qu'une seule frontière
                 possible entre chaque palier. Les 3 nouvelles colonnes
                 « Libellé niveau 0/1/2 » documentent les libellés qualitatifs
                 dans le SQL généré (commentaire), pour faciliter la relecture
                 et la transcription ultérieure en item `select` côté
                 application — elles n'entrent dans aucune condition.

Compatibilité : les types boolean/scale_0_10/text produisent EXACTEMENT le
même SQL qu'avant le 21/08/2026 (mêmes colonnes utilisées, mêmes règles de
génération) — l'extension à 3 niveaux n'introduit aucun changement de
comportement pour les lignes existantes.

Priorité de sécurité : le mécanisme d'évaluation (voir
packages/rules-engine/src/screening/index.ts) applique toujours ROUGE >
ORANGE > VERT, quel que soit le type de la ligne — ce n'est pas une propriété
du gabarit ni de ce script, mais du dispatcher de dépistage lui-même
(règles `medical_referral`/`stop_program` évaluées avant `require_precaution`).

Les règles cliniques combinatoires ou complexes (associations de plusieurs
symptômes, conditions postopératoires, red flags nécessitant plusieurs
variables non couvertes par la colonne « Combiner avec ») ne sont PAS gérées
par ce gabarit : elles continuent à être rédigées directement en SQL, comme
c'était déjà le cas avant cette extension.
"""
import sys
import re
import json
import openpyxl

PATHOLOGY_CODES = {
    "Arthrose genou": "ARTHROSE_GENOU",
    "Arthrose hanche": "ARTHROSE_HANCHE",
    "Polyarthrite rhumatoide": "POLYARTHRITE_RHUMATOIDE",
    "Spondyloarthrite axiale": "SPONDYLOARTHRITE_AXIALE",
}

COLUMNS = [
    "code", "label", "type", "seuil_orange", "seuil_rouge",
    "combiner_avec", "message_orange", "message_rouge", "justification",
    # Colonnes ajoutées le 21/08/2026 (réponse Q1) — uniquement pertinentes
    # pour le type "select_3" ; ignorées pour les autres types (backward
    # compatible : une ligne existante n'ayant pas ces colonnes remplies
    # produit exactement le même SQL qu'avant leur ajout).
    "libelle_vert", "libelle_orange_qualitatif", "libelle_rouge_qualitatif",
]

SELECT_3_DEFAULT_ORANGE = 1
SELECT_3_DEFAULT_ROUGE = 2


def truthy(value):
    if value is None:
        return False
    return str(value).strip().lower() in ("oui", "yes", "true", "1")


def parse_int_or_none(raw):
    raw = str(raw).strip() if raw is not None else ""
    if raw == "":
        return None, None
    try:
        value = float(raw)
    except ValueError:
        return None, f"seuil numérique invalide : {raw!r}"
    if not value.is_integer():
        return None, f"seuil select_3 doit être un entier (0/1/2) : {raw!r}"
    return int(value), None


def base_condition(item_type, code, threshold_raw):
    """Construit la condition élémentaire pour un item, selon son type.

    Pour boolean/scale_0_10/text : comportement STRICTEMENT inchangé depuis
    avant le 21/08/2026 (Q1). Pour select_3 : voir le docstring du module.
    """
    threshold_raw = str(threshold_raw).strip()
    if item_type == "boolean":
        if not truthy(threshold_raw):
            return None, "seuil booléen non reconnu (attendu: Oui)"
        return {"field": code, "operator": "equals", "value": True}, None
    if item_type == "scale_0_10":
        try:
            value = float(threshold_raw)
        except ValueError:
            return None, f"seuil numérique invalide : {threshold_raw!r}"
        if value.is_integer():
            value = int(value)
        return {"field": code, "operator": "gte", "value": value}, None
    if item_type == "select_3":
        value, err = parse_int_or_none(threshold_raw)
        if err:
            return None, err
        if value is None:
            return None, None  # géré par l'appelant (valeur par défaut)
        if value not in (0, 1, 2):
            return None, f"seuil select_3 hors plage (attendu 0, 1 ou 2) : {value!r}"
        return {"field": code, "operator": "gte", "value": value}, None
    # type "text" : pas de condition dérivable automatiquement.
    return None, "type 'text' : condition à écrire manuellement (non générée)"


def select_3_threshold(raw, default):
    """Résout le seuil select_3 effectif (valeur du gabarit, ou valeur par
    défaut 1/2 si la cellule est vide) — voir docstring du module."""
    raw_str = str(raw).strip() if raw is not None else ""
    if raw_str == "":
        return default
    return raw_str


def sql_escape(text):
    return (text or "").replace("'", "''")


def main():
    if len(sys.argv) != 2:
        print("Usage: convert_thresholds_to_rules.py <gabarit_rempli.xlsx>", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    warnings = []
    statements = []

    for sheet_name, pathology in PATHOLOGY_CODES.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]

        # Repère la ligne d'en-tête (contient "Code item") puis saute la
        # ligne d'exemple juste en dessous.
        header_row = None
        for row in ws.iter_rows(min_row=1, max_row=10):
            if any(str(c.value).strip() == "Code item" for c in row if c.value):
                header_row = row[0].row
                break
        if header_row is None:
            warnings.append(f"[{sheet_name}] en-tête introuvable, onglet ignoré")
            continue

        items = {}  # code -> row dict, pour résoudre "combiner_avec"
        rows_data = []
        for r in range(header_row + 2, ws.max_row + 1):  # +2 : saute la ligne d'exemple
            values = {col: ws.cell(row=r, column=i + 1).value for i, col in enumerate(COLUMNS)}
            if not values["code"]:
                continue
            items[values["code"]] = values
            rows_data.append(values)

        for values in rows_data:
            code = values["code"]
            item_type = (values["type"] or "").strip()
            combine_codes = [c.strip() for c in (values["combiner_avec"] or "").split(",") if c.strip()]

            is_select_3 = item_type == "select_3"

            # Cohérence des plages (réponse Q1 : le script doit "valider la
            # cohérence des plages") : pour scale_0_10/select_3, le seuil
            # rouge doit être strictement supérieur au seuil orange lorsque
            # les deux sont fournis explicitement pour le même code (pas de
            # combinaison). Vérifié une seule fois par ligne (pas par
            # niveau) pour ne pas dupliquer l'avertissement.
            if item_type in ("scale_0_10", "select_3") and not combine_codes:
                orange_raw = select_3_threshold(values["seuil_orange"], SELECT_3_DEFAULT_ORANGE) if is_select_3 else values["seuil_orange"]
                rouge_raw = select_3_threshold(values["seuil_rouge"], SELECT_3_DEFAULT_ROUGE) if is_select_3 else values["seuil_rouge"]
                try:
                    if orange_raw not in (None, "") and rouge_raw not in (None, ""):
                        if float(rouge_raw) <= float(orange_raw):
                            warnings.append(
                                f"[{sheet_name}] {code} : seuil rouge ({rouge_raw}) doit être "
                                f"strictement supérieur au seuil orange ({orange_raw}) — ligne(s) "
                                f"générée(s) quand même, à corriger avant relecture"
                            )
                except (TypeError, ValueError):
                    pass

            for level, action, severity, default_select3 in (
                ("seuil_orange", "require_precaution", "warning", SELECT_3_DEFAULT_ORANGE),
                ("seuil_rouge", "medical_referral", "critical", SELECT_3_DEFAULT_ROUGE),
            ):
                raw = values[level]
                if is_select_3:
                    # select_3 génère toujours ses 2 niveaux (valeur par
                    # défaut 1/2 si la cellule est vide) — contrairement à
                    # boolean/scale_0_10/text où une cellule vide signifie
                    # "pas de règle à ce niveau".
                    raw = select_3_threshold(raw, default_select3)
                elif raw is None or str(raw).strip() == "":
                    continue

                own_condition, err = base_condition(item_type, code, raw)
                if err:
                    warnings.append(f"[{sheet_name}] {code} ({level}) : {err}")
                    continue
                if own_condition is None:
                    continue

                sub_conditions = [own_condition]
                skip = False
                for other_code in combine_codes:
                    other = items.get(other_code)
                    if not other:
                        warnings.append(f"[{sheet_name}] {code} ({level}) : code combiné '{other_code}' introuvable")
                        skip = True
                        continue
                    other_raw = other[level]
                    other_type = (other["type"] or "").strip()
                    if other_type == "select_3":
                        other_raw = select_3_threshold(other_raw, default_select3)
                    elif not other_raw:
                        warnings.append(
                            f"[{sheet_name}] {code} ({level}) : '{other_code}' n'a pas de seuil "
                            f"pour ce même niveau, combinaison ignorée"
                        )
                        skip = True
                        continue
                    other_cond, other_err = base_condition(other_type, other_code, other_raw)
                    if other_err:
                        warnings.append(f"[{sheet_name}] {other_code} ({level}) : {other_err}")
                        skip = True
                        continue
                    sub_conditions.append(other_cond)
                if skip:
                    continue

                condition = sub_conditions[0] if len(sub_conditions) == 1 else {"all": sub_conditions}
                message_col = "message_orange" if level == "seuil_orange" else "message_rouge"
                message = values[message_col] or "(message non renseigné dans le gabarit — à compléter)"
                rule_id = f"PROPOSED_{pathology}_{code.upper()}_{'ORANGE' if level == 'seuil_orange' else 'ROUGE'}"
                rule_id = re.sub(r"[^A-Z0-9_]", "_", rule_id)

                condition_json = json.dumps(condition, ensure_ascii=False)
                justification = values["justification"] or ""
                comment_lines = []
                if justification:
                    comment_lines.append(f"-- Source/justification (concepteur médical) : {justification}")
                else:
                    comment_lines.append("-- Aucune justification renseignée dans le gabarit.")
                if is_select_3:
                    lv = values["libelle_vert"] or "(non renseigné)"
                    lo = values["libelle_orange_qualitatif"] or "(non renseigné)"
                    lr = values["libelle_rouge_qualitatif"] or "(non renseigné)"
                    comment_lines.append(
                        f"-- Item qualitatif à 3 niveaux (select_3) — 0=vert: {lv} | "
                        f"1=orange: {lo} | 2=rouge: {lr}"
                    )
                comment = "\n".join(comment_lines)

                statements.append(
                    comment
                    + "\ninsert into public.clinical_rules "
                    "(rule_id, pathology, condition, severity, action, message, active, version) values (\n"
                    f"  '{rule_id}',\n"
                    f"  '{pathology}',\n"
                    f"  '{sql_escape(condition_json)}',\n"
                    f"  '{severity}',\n"
                    f"  '{action}',\n"
                    f"  '{sql_escape(message)}',\n"
                    "  false,\n"
                    "  'V0.1-proposition'\n"
                    ")\non conflict (rule_id) do nothing;\n"
                )

    print("-- Généré automatiquement par convert_thresholds_to_rules.py")
    print("-- TOUTES les règles ci-dessous sont insérées `active = false`.")
    print("-- Relire chaque condition, puis activer explicitement (active = true,")
    print("-- validated_by, validated_date) UNIQUEMENT celles que vous validez.\n")
    for statement in statements:
        print(statement)

    if warnings:
        print("\n-- AVERTISSEMENTS (aucune règle générée pour ces lignes, sauf mention contraire) :", file=sys.stderr)
        for w in warnings:
            print(f"--   {w}", file=sys.stderr)
    print(f"-- {len(statements)} proposition(s) générée(s), {len(warnings)} avertissement(s).", file=sys.stderr)


if __name__ == "__main__":
    main()
